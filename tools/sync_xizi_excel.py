#!/usr/bin/env python3
"""Update XIZI prices from the supplier workbook, preserving form codes and SMEC."""
from __future__ import annotations

import argparse
import hashlib
import json
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from build_pricing_catalog import (
    clean, parse_xizi_base, parse_xizi_containers, parse_xizi_decoration,
    parse_xizi_doors,
)


# Explicit aliases retain saved specification values. Values are supplier sheet codes.
DECORATION_ALIASES = {
    "Car walls": {
        "painted-steel": "Painted Steel", "aisi-443": "AISI443",
        "aisi-304": "AISI304", "ti-gold": "AISI304 GOLD HSS",
    },
    "Floor": {
        "pattern steel": "Рифленая сталь AISI443",
        "Recess in platform 20 mm": "Ниша 20 мм", "NICHE_20": "Ниша 20 мм",
        "Recess in platform 30 mm": "Ниша 30 мм", "NICHE_30": "Ниша 30 мм",
    },
    "Mirror": {"HALF": "Половина высоты", "FULL": "Во всю высоту"},
    "LOP": {f"U-{code}": code for code in ("ZW4300", "ZW4500", "ZW4100", "ZW4200", "ZY500")},
}

OPTION_CELLS = {
    "CONTAINER_40HQ": "B3", "EFS2": "B4", "CWTSAFETY": "B5",
    "ROLLER_GUIDES": "B6", "ISC": "B7", "DHB": "B8", "ADO": "B9",
    "RLEV": "B10", "FAN2": "B11", "COP2": "B12", "DK": "B13",
    "ILED_7": "B14", "ILED_10_4": "B15", "ILED_12_1": "B16",
    "ILED_18_5": "B17", "CCTV": "B18", "IP_CAM": "B19", "CTL": "B20",
    "PKS": "B21", "ARD_15": "B22", "ARD_22": "C22", "ARD_37": "D22",
    "ARD_22_2": "E22", "ARD_37_2": "F22", "TC": "B23",
    "EARTHQUAKE_EMERGENCY_RETURN": "B24", "ACCOLDSMALL": "B25",
    "ACCOLDLARGE": "C25", "ACHEATSMALL": "B26", "ACHEATLARGE": "C26",
    "IC_CARD": "B27", "HAD": "B28",
}

LMR_CODES = {
    3: "RUS_PIT_UNLOCK_DEVICE", 4: "RUS_CONTROLLER_WITH_TRIANGLE",
    5: "VOICE_ANNOUNCEMENT_IN", 6: "RUS_ALL_PLUGS_SHOULD",
    7: "RUS_PIT_EMERGENCY_STOP", 8: "RUS_CAR_TOP_EMERGENCY",
    9: "RUS_ALL_STICKERS_IN", 10: "RUS_ERO_2", 11: "RUS_UPPLER_LOWER",
    12: "RUS_THE_NAMEPLATE_OF", 13: "RUS_EAC_NAMEPLATE_ON",
    14: "RUS_EAC_SYMBOL_AT", 15: "RUS_JUMPER_FOR_SLOW_MOVING",
    16: "RUS_IN_THE_PACKING", 21: "RUS_PIT_INSPECTION_BOX",
    22: "RUS_HOISTWAY_LIGHTING_BY", 23: "RUS_HYDRAULIC_BUFFER_CAPACITY",
}

# Keep existing quantity rules; the Excel supplies their base/unit prices.
# P denotes Price from the cited supplier cell. R/K/S are dimensions in metres.
FORMULAS = {
    "CWTSAFETY": "P+(R+K+S)*4", "COP2": "P+(N-10)*45",
    "CCTV": "P*(R+K+16)", "TC": "P*(R+K+16)",
    "EARTHQUAKE_EMERGENCY_RETURN": "P+300*((R+K+S)/1.5-(R+K+S)/2.5)",
    **{code: "P+11*(R+K+12)" for code in (
        "ACCOLDSMALL", "ACCOLDLARGE", "ACHEATSMALL", "ACHEATLARGE")},
    "HAD": "P+101*D", "RUS_PIT_INSPECTION_BOX": "P+24*(R+K+16)",
    "RUS_HOISTWAY_LIGHTING_BY": "P*((R+K+S)/4-(R+K+S)/7)",
}


def unique_entries(entries, fields):
    result = {}
    for entry in entries:
        key = tuple(entry[field] for field in fields)
        if key in result and result[key] != entry:
            raise ValueError(f"Conflicting supplier entries: {key}")
        result[key] = entry
    return list(result.values())


PRICE_KEYS = {
    "basePrices": ("series", "capacity", "speed", "stops"),
    "doors": ("manufacturer", "part", "doorType", "fireRating", "finish", "capacity", "floor", "width"),
    "decorations": ("category", "code"),
    "options": ("code",),
}


def apply_supplement(catalog, supplement):
    """Apply the stored source priority; default to filling missing prices only."""
    xizi = catalog["xizi"]
    prefer_reference = supplement.get("priority") == "reference-project"
    for group, fields in PRICE_KEYS.items():
        for addition in supplement.get(group, []):
            is_requirement = group == "options" and (addition.get("category") == "russia"
                or any(e.get("code") == addition.get("code") for e in xizi["localRequirements"]))
            target = "localRequirements" if is_requirement else group
            entries = xizi[target]
            if target != group:
                # Preserve automatic requirements, including voice announcements.
                xizi[group] = [e for e in xizi[group] if e.get("code") != addition.get("code")]
            index = {tuple(entry.get(field) for field in fields): entry for entry in entries}
            key = tuple(addition.get(field) for field in fields)
            current = index.get(key)
            if current is None:
                current = deepcopy(addition)
                entries.append(current)
                index[key] = current
            elif prefer_reference or current.get("price") is None or current.get("source", {}).get("kind") == "reference-project":
                current.update(deepcopy(addition))
    xizi["priceSource"]["supplement"] = deepcopy(supplement["source"])
    xizi["priceSource"]["priority"] = "reference-project" if prefer_reference else "excel"
    xizi["priceSource"]["seriesWithoutPrices"] = [
        name for name in xizi["series"] if not any(e["series"] == name for e in xizi["basePrices"])
    ]
    return catalog


def sync_catalog(catalog, workbook, source_name, source_hash):
    result = deepcopy(catalog)
    xizi = result["xizi"]
    base = unique_entries(parse_xizi_base(workbook), ("series", "capacity", "speed", "stops"))
    if not base:
        raise ValueError("No XIZI base prices found")
    xizi["basePrices"] = base
    # Keep unsupported model names visible so saved specifications fail explicitly.
    xizi["series"] = list(dict.fromkeys([*xizi["series"], *(e["series"] for e in base)]))
    xizi["doors"] = unique_entries(parse_xizi_doors(workbook), (
        "manufacturer", "part", "doorType", "fireRating", "finish", "capacity", "floor", "width"))
    xizi["containers"] = parse_xizi_containers(workbook)

    source_decorations = {(e["category"], e["code"]): e for e in parse_xizi_decoration(workbook)}
    covered = set()
    for entry in xizi["decorations"]:
        category, code = entry["category"], entry["code"]
        source_code = DECORATION_ALIASES.get(category, {}).get(code, code)
        key = (category, source_code)
        source = source_decorations.get(key)
        if source is not None:
            entry.update({field: source.get(field) for field in ("price", "overprice", "height")})
            entry["source"] = {"sheet": "Decoration", "code": source_code}
            covered.add(key)
        elif (category, code) != ("Mirror", "NO"):
            entry.update(price=None, overprice=None, height=None)
            entry["source"] = {"status": "not-in-workbook"}
    for key, source in source_decorations.items():
        if key not in covered:
            xizi["decorations"].append({**source, "source": {"sheet": "Decoration", "code": source["code"]}})

    existing_options = {e["code"]: e for e in xizi["options"]}
    existing_lmr = {e["code"]: e for e in xizi["localRequirements"]}
    xizi["options"] = []
    for code in dict.fromkeys([*existing_options, *OPTION_CELLS]):
        if code in LMR_CODES.values():
            continue  # Included once, as an automatic local requirement.
        entry = existing_options.get(code, {"code": code, "category": "Options"})
        cell = OPTION_CELLS.get(code)
        entry.update(price=clean(workbook["Options"][cell].value) if cell else None,
                     type="formula" if code in FORMULAS else "fixed",
                     formula=FORMULAS.get(code))
        entry.pop("prices", None)
        entry["source"] = {"sheet": "Options", "cell": cell} if cell else {"status": "not-in-workbook"}
        xizi["options"].append(entry)

    xizi["localRequirements"] = []
    for row, code in LMR_CODES.items():
        entry = existing_lmr.get(code, existing_options.get(code, {"code": code}))
        entry.update(category="russia", price=clean(workbook["LMR"].cell(row, 2).value),
                     type="formula" if code in FORMULAS else "fixed", formula=FORMULAS.get(code),
                     source={"sheet": "LMR", "cell": f"B{row}"})
        entry.setdefault("description", clean(workbook["LMR"].cell(row, 1).value))
        xizi["localRequirements"].append(entry)

    xizi["priceSource"] = {
        "kind": "excel", "file": source_name, "sha256": source_hash,
        "quantityRules": "Existing calculator rules; P is the base/unit price from the supplier workbook.",
        "seriesWithoutPrices": [name for name in xizi["series"] if not any(e["series"] == name for e in base)],
    }
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--catalog", type=Path, default=Path("src/TFlexDrawingService.Api/Data/pricing-catalog.json"))
    parser.add_argument("--supplement", type=Path, help="Default: xizi-price-supplement.json beside the catalog")
    parser.add_argument("--check", action="store_true", help="Verify catalog prices without writing files")
    args = parser.parse_args()
    catalog = json.loads(args.catalog.read_text())
    workbook = load_workbook(args.workbook, data_only=True)
    result = sync_catalog(catalog, workbook, args.workbook.name, hashlib.sha256(args.workbook.read_bytes()).hexdigest())
    workbook.close()
    supplement_path = args.supplement or args.catalog.with_name("xizi-price-supplement.json")
    if supplement_path.exists():
        apply_supplement(result, json.loads(supplement_path.read_text()))
    if args.check:
        if result != catalog:
            raise SystemExit("XIZI catalog differs from the workbook plus stored reference priority; run without --check to update")
    else:
        result["generatedAt"] = datetime.now(timezone.utc).isoformat()
        args.catalog.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    xizi = result["xizi"]
    print(json.dumps({
        "checked" if args.check else "updated": {key: len(xizi[key]) for key in (
            "basePrices", "doors", "decorations", "options", "localRequirements", "containers")},
        "seriesWithoutPrices": xizi["priceSource"]["seriesWithoutPrices"],
        "unpricedOptions": [e["code"] for e in xizi["options"] if e["price"] is None],
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
