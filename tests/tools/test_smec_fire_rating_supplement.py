import sys
import unittest
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools"))
import sync_smec_reference
import build_pricing_catalog


class SmecFireRatingSupplementTests(unittest.TestCase):
    def supplement(self, code="Firerated door(EI60)", capacity=630, variant="2S"):
        return {
            "category": "DoorAddon", "code": code, "capacity": capacity,
            "variant": variant, "price": 123,
            "source": {
                "kind": "excel-supplement",
                "file": sync_smec_reference.FIRE_DOOR_SOURCE_FILE,
                "sheet": "Decoration", "cell": "D120",
            },
        }

    def test_preserves_only_missing_marked_fire_doors_once_and_primary_wins(self):
        supplement = self.supplement()
        unrelated = self.supplement("old door")
        unrelated["source"] = dict(supplement["source"])
        unmarked = self.supplement("Firerated door(E60)")
        unmarked.pop("source")
        for primary_price in (0, -1):
            primary = [{**self.supplement(), "price": primary_price}]
            result = sync_smec_reference.preserve_fire_door_supplement(
                primary, [supplement, dict(supplement), self.supplement("Firerated door(E60)"), unrelated, unmarked]
            )
            self.assertEqual([primary_price], [item["price"] for item in result if item["code"] == supplement["code"]])
            self.assertEqual(1, sum(item["code"] == "Firerated door(E60)" for item in result))
            self.assertFalse(any(item["code"] == "old door" for item in result))

    def test_catalog_parser_keeps_original_fire_rating_codes(self):
        workbook = Workbook()
        ws = workbook.active
        ws.title = "Decoration"
        ws.cell(105, 4, 630)
        ws.cell(120, 2, "Firerated door(E60)")
        ws.cell(120, 4, 10)
        ws.cell(121, 2, "Firerated door(EI60)")
        ws.cell(121, 4, 20)
        entries = build_pricing_catalog.parse_smec_decorations(workbook)
        fire_doors = [item for item in entries if item["category"] == "DoorAddon"]
        self.assertEqual(["Firerated door(E60)", "Firerated door(EI60)"], [item["code"] for item in fire_doors])
        self.assertEqual(
            [
                {"kind": "excel-supplement", "file": build_pricing_catalog.LEHY_PRICE_SOURCE.name, "sheet": "Decoration", "cell": "D120"},
                {"kind": "excel-supplement", "file": build_pricing_catalog.LEHY_PRICE_SOURCE.name, "sheet": "Decoration", "cell": "D121"},
            ],
            [item["source"] for item in fire_doors],
        )


if __name__ == "__main__":
    unittest.main()
